import re
import pandas as pd
from pandas import io
import numpy as np
import pdfplumber
from openpyxl.styles import Alignment, PatternFill, Font


filename = "WEB-•-REVIEW-OF-THE-ECONOMY-2025.pdf"
LINE_TOL = 3.0
WRAP_TOL = 7.0
COL_GAP = 20.0
SUPERSCRIPT = 0.75
HEADER_FILL = PatternFill("solid", start_color="D9E1F2")   # light blue
BOLD = Font(bold=True)

excel_label_map = {
        188: "GDP Const 2012 Prices TTD Mn",
        189: "GDP Const 2012 Prices Pct Chg",
        190: "GDP Const 2012 Prices Pct Cont",
        191: "GDP Curr Prices TTD Mn",
        192: "GDP Curr Prices Pct Chg",
        193: "GDP Curr Prices Pct Cont"
    }

suffix_map = {
                "e": "(MOF Estimated)",
                "r": "(CSO Revised)",
                "p": "(CSO Provisional)"
            }

notes_page = pd.DataFrame(
    [
        ("Source", "Review of the Economy 2025 (Ministry of Finance, Trinidad and Tobago), GDP appendices 1-6, PDF pages 189-194."),
        ("Sheets", "Six tables: GDP by industry at constant 2012 prices and at current prices. Each has a TT$ Mn sheet, a Pct Chg sheet and a Pct Cont sheet."),
        ("Units", "TT$ Mn sheets are in TT$ millions. Pct Chg and Pct Cont sheets are percentages."),
        ("Column labels", "ISIC = International Standard of Industrial Classification of All Economic Activities. CSO Revised / CSO Provisional = Central Statistical Office figures. MOF Estimated = Ministry of Finance estimate. Q4* and Q1* are quarterly columns; the asterisk means that Q4 is referred to the period October to December while Q1 refers to the period January to March."),
        ("Pct Chg", "Annual columns are year-on-year % change. Checked by recomputing from the TT$ Mn sheets for 2022-2024."),
        ("Pct Cont", "Share of GDP. Annual columns: GDP at purchaser prices = 100. Quarterly columns: GDP at basic prices = 100, because the source gives no purchaser-price total for quarters."),
        ("Blank cells", "The PDF shows '-' or 'N/A' in these cells. Both are left empty here, so a blank means no figure was given in the source."),
        ("Extraction", "Automated with Python (pdfplumber and pandas): words are grouped into rows by position and each number is placed in a column by its right edge. Wrapped row labels are merged, footnote markers removed."),
    ],
    columns=["Item", "Detail"],
)

def drop_superscripts(word_bank: pd.DataFrame) -> pd.DataFrame | None:
    keep_rows = []
    if word_bank is None:
        return None
    else:
        for _, row in word_bank.iterrows():
            same_line = word_bank[(word_bank["top"] - row["top"]).abs() <= 4]
            biggest_font_size = same_line["size"].max()
            # print(type(biggest_font_size))
            # print(biggest_font_size)
            keep_rows.append(row["size"] >= (biggest_font_size * SUPERSCRIPT))
        return word_bank[keep_rows]

def get_right_column(x: float, columns: list[tuple[float, float]]) -> int | None:
    if columns is None:
        return None
    else:
        bounds = [(columns[i][1] + columns[i + 1][0]) / 2 for i in range(len(columns)-1)]
        return int(np.searchsorted(bounds, x))

def convert_to_number(cell: str):
    cell = cell.strip()

    if cell in ("", "-", "N/A"):
        return np.nan
    negative_num = cell.startswith("(") and cell.endswith(")")
    try:
        value = float(cell.strip("()").replace(",", ""))
    except ValueError:
        return np.nan
    return -value if negative_num and value != 0 else value

def join_header_columns(header_words: pd.DataFrame, gap: float = 8.0) -> list[str] | None:
    #Join header words into required cells
    if header_words is None:
        return None
    else:

        cells, current, last_x1 = [], [], None
        for _, word in header_words.sort_values("x0").iterrows():
            if last_x1 is not None and word["x0"] - last_x1 > gap:
                cells.append(current)
                current = []
            glue = bool(current) and word["x0"] - last_x1 < 1
            current.append(("" if glue else " ") + word["text"])
            last_x1 = word["x1"]
        if current:
            cells.append(current)
        names = ["".join(c).strip() for c in cells]
        return [re.sub(r"(?<=[A-Za-z])\d$", "", n) for n in names]


def find_data_columns(data_df: pd.DataFrame) -> list[tuple[float, float]] | None:
    if data_df is None:
        return None
    else:
        sorted_x1 = np.sort(data_df["x1"].to_numpy())
        columns, start = [], 0
        for i in range(1, len(sorted_x1)):
            if (sorted_x1[i] - sorted_x1[i - 1]) > COL_GAP:
                columns.append((sorted_x1[start], sorted_x1[i - 1]))
                start = i
        columns.append((sorted_x1[start], sorted_x1[-1]))
        return columns


def extract_word_bank(filename: str, page_index: int) -> pd.DataFrame | None:
    if filename.endswith(".pdf"):
        try:
            with pdfplumber.open(filename) as pdf:
                sample_page = pdf.pages[page_index]
                word_bank = sample_page.extract_words(extra_attrs = ["size"])
                word_bank_df = pd.DataFrame(word_bank)
                word_bank_df = word_bank_df[~word_bank_df["text"].str.startswith("(cid:")]
                return word_bank_df
                
        except:
            if FileExistsError or FileNotFoundError:
                print("File Not Found!")
                return None
    

    else:
        print("Incorrect file passed!")
        return None

def row_aggregator(words_df: pd.DataFrame, tolerance: float = LINE_TOL) -> list[pd.DataFrame] | None:
    try:
        if words_df.empty:
            return None

        words_df = words_df.sort_values(by= "top")
        row_id = (words_df["top"].diff() > tolerance).cumsum()
        sorted_df_list = []
        for _, g in words_df.groupby(row_id):
            
            sorted_df_list.append(g.sort_values("x0"))
        return sorted_df_list
        
    except Exception as e:
        print(f"Row aggregation error {e}")
        return None

def parse_gdp_table(filename: str, page_index: int) -> pd.DataFrame | None:
    word_bank_df = extract_word_bank(filename= filename, page_index= page_index)
    

    if word_bank_df is None:
        return None
    else:
        rows_list = row_aggregator(words_df= word_bank_df)


        #Define boundaries of table
        header_top = word_bank_df.loc[word_bank_df["text"] == "INDUSTRY", "top"].min()
        header_btm = word_bank_df.loc[word_bank_df["text"] == "INDUSTRY", "bottom"].min()
        source_top = word_bank_df.loc[word_bank_df["text"].str.startswith("Source"), "top"].min()

        #Define title
        if rows_list is None:
            return None
        else:
            title_lines = []
            above_header_lines = row_aggregator(words_df= word_bank_df[word_bank_df["top"] < header_top - 4])

            if above_header_lines is None:
                return None
            else:
                for line in above_header_lines:
                    line_text = " ".join(line["text"])
                    if "Appendix" in line_text or "Gross" in line_text:
                        title_lines.append(line_text)
                title = " ".join(title_lines)

            # Extract header text
            header = word_bank_df[(word_bank_df["top"] >= header_top - 4) & (word_bank_df["top"] < header_btm + 2)]

            #Extract and clean up body text

            body = drop_superscripts(word_bank= word_bank_df[(word_bank_df["top"] >= header_btm + 2) & (word_bank_df["top"] < source_top - 3)])

            # Split data between text and data by measured margin

            split_x_line = header.loc[header["text"].str.startswith("ISIC"), "x0"].min() - 3
            if body is None:
                return None
            else:
                label_words = body[body["x0"] < split_x_line]
                data_words = body[body["x0"] >= split_x_line]

                # Split data_words table into columns

                columns = find_data_columns(data_df= data_words)

                names = join_header_columns(header_words=header[header["x0"] >= split_x_line])

                if columns is None or names is None:
                    return None
                elif len(columns) != len(names):
                    print(f"WARNING page {page_index}: {len(names)} header cells but {len(columns)} data columns")
                    names = [f"col{i}" for i in range(len(columns))]
                else:
                    cleaned_names = []
                    for name in names:
                        name.strip()
                        last_char = name[-1].lower()
                        if last_char in suffix_map:
                            cleaned_name = name[:-1] + " " + suffix_map[last_char]
                        else:
                            cleaned_name = name
                        cleaned_names.append(cleaned_name)
                    names = cleaned_names

                # --- data rows: each line of data words becomes one row, numbers placed by x1 ---

                data_rows = []
                raw_data_rows = row_aggregator(words_df=data_words)
                if raw_data_rows is None:
                    return None
                else:
                    for line in raw_data_rows:
                        cells = [""] * len(columns)
                        for _, word in line.iterrows():
                            right_col = get_right_column(x= word["x1"], columns= columns)
                            if right_col is None:
                                return None
                            else:
                                cells[right_col] = word["text"]
                        data_rows.append({"y": line["top"].mean(), "cells": cells, "labels": []})

                # --- label lines: attach each to the nearest data row, so wrapped labels stay together ---
                label_rows = []
                raw_label_rows = row_aggregator(label_words)
                if raw_label_rows is None:
                    return None
                else:
                    for line in raw_label_rows:
                        y_coord, text = line["top"].mean(), " ".join(line["text"])
                        nearest = None
                        smallest_dist = float("inf")
                        for row in data_rows:
                            distance = abs(row["y"] - y_coord)
                            if distance < smallest_dist:
                                smallest_dist = distance
                                nearest = row
                        if nearest is not None and abs(nearest["y"] - y_coord) <= WRAP_TOL:
                            nearest["labels"].append((y_coord, text))
                        else:
                            label_rows.append({"y": y_coord, "cells": [""] * len(columns), "labels": [(y_coord, text)]})
                full_rows = sorted(label_rows + data_rows, key=lambda row: row["y"])
                records = []
                for row in full_rows:
                    label_part = sorted(row["labels"])
                    label_text = [text for y, text in label_part]
                    full_label = " ".join(label_text)
                    records.append([full_label] + row["cells"])

                df = pd.DataFrame(records, columns=["INDUSTRY"] + names)
                for col in names:
                    if col != "ISIC":
                        df[col] = df[col].map(convert_to_number)
                df.attrs["title"] = title

                return df




            

if __name__ == "__main__":

    with pd.ExcelWriter("GDP_Analysis_Output.xlsx", engine="openpyxl") as writer:
        for page_index in range(188, 194):
            page_df = parse_gdp_table(filename=filename, page_index=page_index)
            if page_df is not None:
                
                # print(page_df.to_string(index=False))
                # print(page_df.attrs["title"])

                sheet_tab_name = excel_label_map.get(page_index, f"Appendix_Page_{page_index + 1}")
                page_df.to_excel(writer, sheet_name=sheet_tab_name, index=False)
                if "TTD Mn" in sheet_tab_name:
                    num_format = "#,##0.0"
                else:
                    num_format = "0.0"
                active_sheet = writer.sheets[sheet_tab_name]
                for row in active_sheet.iter_rows(min_row= 2, min_col= 3):
                    for cell in row:
                        cell.number_format = num_format

                active_sheet.column_dimensions["A"].width = 50
                active_sheet.column_dimensions["B"].width = 9
                for col_letter in "CDEFGHI":
                    active_sheet.column_dimensions[col_letter].width = 15

                for cell in active_sheet[1]:
                    cell.font = BOLD
                    cell.fill = HEADER_FILL
                    cell.alignment = Alignment(wrap_text= True, horizontal= "center", vertical= "center")
                active_sheet.row_dimensions[1].height = 32

                for row in active_sheet.iter_rows(min_row= 2):
                    label = str(row[0].value or "")
                    if label.startswith(("GDP AT", "Memo Items")):
                        for cell in row:
                            cell.font = BOLD

                active_sheet.freeze_panes = "C2"

        notes_page.to_excel(writer, sheet_name="Notes", index=False)
        sheet = writer.sheets["Notes"]
        sheet.column_dimensions["A"].width = 16
        sheet.column_dimensions["B"].width = 120
        for row in sheet.iter_rows(min_row = 2):
            for cell in row:
                cell.alignment = Alignment(wrap_text= True, vertical="top")